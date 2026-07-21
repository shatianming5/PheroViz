"""Deterministic, fail-closed experiment-case proposals for curated candidates."""

from __future__ import annotations

from copy import deepcopy
import csv
from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import subprocess
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import pandas as pd
except ImportError:
    pd = None

from .provenance import sha256_file


SCHEMA_VERSION = "1.0"
EVALUATION_SCHEMA_VERSION = "1.1.0"
DEFAULT_MAX_FILE_BYTES = 64 * 1024 * 1024
DEFAULT_MAX_ROWS = 100_000
DEFAULT_MAX_COLUMNS = 64
PROPOSAL_RULE_V1 = "simple-2d-v1"
PROPOSAL_RULE_V2 = "simple-2d-v2"
PROPOSAL_RULE_V3 = "simple-2d-v3"
PROPOSAL_RULE_V4 = "simple-2d-v4"
CURRENT_PROPOSAL_RULE = PROPOSAL_RULE_V2
WIDE_MELT_BINDING_MODE = "wide_melt"
WIDE_MELT_GROUP_COLUMN = "__wide_group__"
WIDE_MELT_VALUE_COLUMN = "__wide_value__"
GROUPED_DOT_PLOT_MIN_VALUES = 5
MAX_BAR_CATEGORICAL_X = 200
RENDERABILITY_POLICY_ID = "outcome-independent-static-renderability-v1"
RENDERABILITY_POLICY_HASH = (
    "283f0520bf1fda98773958602be6206c1776c573a00273391dbfeb6acd565cec"
)
CONTROLLED_NUMERIC_COLUMN_PATTERN = re.compile(
    r"(?i)(?:^|[^a-z])"
    r"(?:dose|dosage|concentration|conc|wavelength|frequency|freq|"
    r"voltage|current|temperature|temp|position)"
    r"(?:$|[^a-z])"
)
SCATTERING_VECTOR_COLUMN_PATTERN = re.compile(
    r"(?i)^\s*q\s*[\(\[]\s*(?:å|a|angstrom)\s*[-−–]?\s*1\s*[\)\]]\s*$"
)
DIRECT_GROUP_RESPONSE_PATTERN = re.compile(
    r"(?i)(?:^|[^a-z])(?:number|size|ph)(?:$|[^a-z])"
)
V4_TEMPORAL_NAME_PATTERN = re.compile(
    r"(?i)(?:^|[^a-z])"
    r"(?:time|timing|duration|elapsed|latency)"
    r"(?:$|[^a-z])"
)
TEMPORAL_UNIT_NAMES = frozenset(
    {
        "ms",
        "msec",
        "s",
        "sec",
        "second",
        "seconds",
        "min",
        "minute",
        "minutes",
        "h",
        "hr",
        "hour",
        "hours",
        "day",
        "days",
        "week",
        "weeks",
        "month",
        "months",
        "year",
        "years",
    }
)
MEASUREMENT_COLUMN_PATTERN = re.compile(
    r"(?i)(?:^|[^a-z])"
    r"(?:value|signal|response|measurement|mean|median|average|"
    r"intensity|count|score|density|level|ratio|rate|length|distance|"
    r"area|volume|mass|weight|height|width|diameter|angle|depth|"
    r"fraction|parameter|data)"
    r"(?:$|[^a-z])"
)
INDEX_HEADER_TOKENS = frozenset(
    {
        "#",
        "n",
        "no",
        "no.",
        "nr",
        "nr.",
        "num",
        "num.",
        "id",
        "idx",
        "index",
        "number",
        "serial",
        "order",
        "obs",
        "rank",
        "row",
        "count",
    }
)
INDEX_HEADER_WORDS = (
    "animal",
    "mouse",
    "mice",
    "rat",
    "sample",
    "subject",
    "replicate",
    "rep",
    "cell",
    "specimen",
    "patient",
    "donor",
    "individual",
    "fish",
    "worm",
    "embryo",
    "larva",
    "fly",
    "well",
    "trial",
)
RENDERABILITY_RULE_ID = "bar-column-categorical-x-cardinality"
BAR_CARDINALITY_REJECTION = "bar_categorical_x_exceeds_200"
RENDERABILITY_POLICY: dict[str, Any] = {
    "schema_version": SCHEMA_VERSION,
    "policy_id": RENDERABILITY_POLICY_ID,
    "declared_inputs": [
        "source_table",
        "source_sheet",
        "evaluation_expectation",
    ],
    "forbidden_inputs": [
        "run_records",
        "method_names",
        "method_scores",
        "failed_case_ids",
        "production_artifacts",
    ],
    "split_scope": ["train", "val", "test"],
    "panel_scope": "every_panel",
    "multi_panel_rule": "reject_case_if_any_panel_fails",
    "rules": [
        {
            "rule_id": RENDERABILITY_RULE_ID,
            "expectation_series_kinds": ["bar"],
            "semantic_families": ["bar", "column"],
            "max_unique_categorical_x_per_panel": MAX_BAR_CATEGORICAL_X,
            "comparison": "less_than_or_equal",
            "x_counting": (
                "union_of_resolved_bar_series_x_after_expectation_filters"
            ),
            "missing_x_values": "count_as_one_category",
        }
    ],
    "audit_counting": {
        "series_count": "all_expected_series_in_panel",
        "expected_points": (
            "sum_per_series_max_resolved_x_y_value_length"
        ),
    },
    "explicitly_uncapped_series_kinds": ["line", "scatter"],
    "non_bar_cardinality_action": "no_cap",
}
RUN_ORDER_COLUMN_PATTERN = re.compile(
    r"run[\s_.-]*order",
    flags=re.I,
)
TEMPORAL_COLUMN_PATTERN = re.compile(
    r"(?i)(?:^|[^a-z])"
    r"(?:time|timepoint|date|datetime|timestamp|year|month|day)"
    r"(?:$|[^a-z])"
)
PAREN_UNIT_PATTERN = re.compile(r"\(([^()]{1,32})\)\s*$")
BRACKET_UNIT_PATTERN = re.compile(r"\[([^\[\]]{1,32})\]\s*$")
SUFFIX_UNIT_PATTERN = re.compile(
    r"(?i)(?:_|-)"
    r"(ms|s|min|h|hz|kg|g|mg|ug|µg|m|cm|mm|um|µm|nm|l|ml|mol|mmol|%)$"
)


class ProposalRejected(ValueError):
    """A deterministic rejection with one or more machine-readable reasons."""

    def __init__(
        self,
        *reasons: str,
        detail: str | None = None,
        audit: dict[str, Any] | None = None,
    ) -> None:
        self.reasons = tuple(sorted(set(reasons or ("proposal-rejected",))))
        self.detail = detail
        self.audit = deepcopy(audit) if audit is not None else None
        super().__init__(",".join(self.reasons))


def _resolve_proposal_rule_version(proposal: dict[str, Any]) -> str:
    value = proposal.get("proposal_rule_version")
    if value is None:
        return PROPOSAL_RULE_V1
    if not isinstance(value, str) or value not in {
        PROPOSAL_RULE_V1,
        PROPOSAL_RULE_V2,
        PROPOSAL_RULE_V3,
        PROPOSAL_RULE_V4,
    }:
        raise ValueError("proposal-rule-version-unsupported")
    return value


@dataclass(frozen=True)
class ColumnProfile:
    name: str
    numeric: bool
    explicit_temporal: bool
    temporal_valid: bool
    categorical: bool
    monotonic_numeric: bool
    non_null_count: int
    unique_count: int


@dataclass(frozen=True)
class TableAnalysis:
    x: str
    y: tuple[str, ...]
    x_mode: str
    chart_family: str
    units: dict[str, str]
    rows: int
    columns: int
    non_null_y: dict[str, int]
    binding_mode: str = "direct"
    wide_value_columns: tuple[str, ...] = ()
    dropped_index_columns: tuple[str, ...] = ()


def _canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )


def _validate_renderability_policy() -> None:
    actual = hashlib.sha256(
        _canonical_json(RENDERABILITY_POLICY).encode("utf-8")
    ).hexdigest()
    if (
        actual != RENDERABILITY_POLICY_HASH
        or RENDERABILITY_POLICY["policy_id"] != RENDERABILITY_POLICY_ID
        or RENDERABILITY_POLICY["rules"][0][
            "max_unique_categorical_x_per_panel"
        ]
        != MAX_BAR_CATEGORICAL_X
    ):
        raise RuntimeError("renderability-policy-mutated")


def _value_identity(value: Any) -> str:
    try:
        missing = bool(pd.isna(value))
    except (TypeError, ValueError):
        missing = False
    if missing:
        return "missing:null"
    if hasattr(value, "item"):
        try:
            value = value.item()
        except (TypeError, ValueError):
            pass
    try:
        encoded = _canonical_json(value)
    except (TypeError, ValueError):
        encoded = _canonical_json(str(value))
    return f"{type(value).__name__}:{encoded}"


def _renderability_audit(
    frame: Any,
    analysis: TableAnalysis,
) -> dict[str, Any]:
    _validate_renderability_policy()
    bar_series_count = (
        len(analysis.y) if analysis.chart_family == "bar" else 0
    )
    identities: set[str] = set()
    if bar_series_count:
        if analysis.binding_mode == WIDE_MELT_BINDING_MODE:
            identities.update(
                _value_identity(value)
                for value in analysis.wide_value_columns
            )
        else:
            resolved_x = frame[analysis.x].tolist()
            for _ in analysis.y:
                identities.update(_value_identity(value) for value in resolved_x)
        unique_categorical_x: int | None = len(identities)
        if unique_categorical_x > MAX_BAR_CATEGORICAL_X:
            decision = "rejected"
            reason = BAR_CARDINALITY_REJECTION
        else:
            decision = "accepted"
            reason = "bar_categorical_x_at_or_below_200"
    else:
        unique_categorical_x = None
        decision = "accepted"
        reason = "non_bar_or_column_no_cardinality_cap"
    return {
        "schema_version": SCHEMA_VERSION,
        "policy_id": RENDERABILITY_POLICY_ID,
        "policy_hash": RENDERABILITY_POLICY_HASH,
        "rule_id": RENDERABILITY_RULE_ID,
        "decision": decision,
        "reason": reason,
        "chart_family": analysis.chart_family,
        "bar_series_count": bar_series_count,
        "unique_categorical_x": unique_categorical_x,
        "max_unique_categorical_x_per_panel": MAX_BAR_CATEGORICAL_X,
        "x_counting": (
            "union_of_resolved_bar_series_x_after_expectation_filters"
        ),
        "missing_x_values": "count_as_one_category",
    }


def _candidate_group_key(
    value: dict[str, Any],
) -> tuple[str, int] | None:
    doi = str(value.get("doi") or "")
    figure_no = value.get("figure_no")
    if not doi or not isinstance(figure_no, int):
        return None
    return doi, figure_no


def _require_dependencies() -> None:
    missing = []
    if pd is None:
        missing.append("pandas")
    if load_workbook is None:
        missing.append("openpyxl")
    if missing:
        raise ProposalRejected(
            "proposal-dependency-missing",
            detail=",".join(missing),
        )


def _validate_headers(values: Iterable[Any]) -> list[str]:
    headers: list[str] = []
    normalized: list[str] = []
    for value in values:
        if value is None or not str(value).strip():
            raise ProposalRejected("column-name-missing")
        name = str(value)
        headers.append(name)
        normalized.append(name.strip().casefold())
    if len(set(normalized)) != len(normalized):
        raise ProposalRejected("column-name-duplicate")
    return headers


def _read_csv(
    path: Path,
    *,
    max_rows: int,
    max_columns: int,
) -> Any:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            header = next(reader)
    except StopIteration as exc:
        raise ProposalRejected("table-empty") from exc
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ProposalRejected(
            "csv-read-error",
            detail=f"{type(exc).__name__}:{exc}",
        ) from exc
    if len(header) > max_columns:
        raise ProposalRejected("table-column-limit")
    _validate_headers(header)
    try:
        frame = pd.read_csv(path, nrows=max_rows + 1)
    except Exception as exc:
        raise ProposalRejected(
            "csv-read-error",
            detail=f"{type(exc).__name__}:{exc}",
        ) from exc
    if len(frame.index) > max_rows:
        raise ProposalRejected("table-row-limit")
    if len(frame.columns) > max_columns:
        raise ProposalRejected("table-column-limit")
    return frame


def _read_xlsx(
    path: Path,
    *,
    sheet_name: str | None,
    max_rows: int,
    max_columns: int,
) -> Any:
    if not sheet_name:
        raise ProposalRejected("xlsx-sheet-required")
    workbook = None
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        if sheet_name not in workbook.sheetnames:
            raise ProposalRejected("xlsx-sheet-missing")
        worksheet = workbook[sheet_name]
        row_iter = worksheet.iter_rows(
            min_row=1,
            max_col=max_columns + 1,
            values_only=True,
        )
        try:
            raw_header = next(row_iter)
        except StopIteration as exc:
            raise ProposalRejected("table-empty") from exc
        if len(raw_header) > max_columns and any(
            value is not None for value in raw_header[max_columns:]
        ):
            raise ProposalRejected("table-column-limit")
        header_values = list(raw_header[:max_columns])
        while header_values and header_values[-1] is None:
            header_values.pop()
        headers = _validate_headers(header_values)
        rows: list[list[Any]] = []
        for row_index, raw_row in enumerate(row_iter, 1):
            if row_index > max_rows:
                raise ProposalRejected("table-row-limit")
            if len(raw_row) > len(headers) and any(
                value is not None for value in raw_row[len(headers) :]
            ):
                raise ProposalRejected("table-ragged-extra-columns")
            rows.append(list(raw_row[: len(headers)]))
        return pd.DataFrame(rows, columns=headers)
    except ProposalRejected:
        raise
    except Exception as exc:
        raise ProposalRejected(
            "xlsx-read-error",
            detail=f"{type(exc).__name__}:{exc}",
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def read_candidate_table(
    candidate: dict[str, Any],
    *,
    max_file_bytes: int = DEFAULT_MAX_FILE_BYTES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
) -> Any:
    _require_dependencies()
    source = candidate.get("source_table") or {}
    path_value = source.get("path")
    if not path_value:
        raise ProposalRejected("source-table-path-missing")
    path = Path(str(path_value))
    if path.is_symlink():
        raise ProposalRejected("source-table-symlink")
    if not path.is_file():
        raise ProposalRejected("source-table-missing")
    size = path.stat().st_size
    if size > max_file_bytes:
        raise ProposalRejected("source-table-file-size-limit")
    if source.get("size_bytes") is not None and int(source["size_bytes"]) != size:
        raise ProposalRejected("source-table-size-mismatch")
    expected_hash = source.get("sha256")
    if not expected_hash or sha256_file(path) != expected_hash:
        raise ProposalRejected("source-table-checksum-mismatch")
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        frame = _read_csv(
            path,
            max_rows=max_rows,
            max_columns=max_columns,
        )
    elif suffix == ".xlsx":
        frame = _read_xlsx(
            path,
            sheet_name=source.get("sheet_name"),
            max_rows=max_rows,
            max_columns=max_columns,
        )
    else:
        raise ProposalRejected("source-table-format-unsupported")
    if len(frame.columns) < 2 or len(frame.columns) > 6:
        raise ProposalRejected("table-column-count-not-simple")
    frame = frame.replace(r"^\s*$", pd.NA, regex=True).dropna(how="all")
    if frame.empty:
        raise ProposalRejected("table-empty")
    empty_columns = [
        str(column) for column in frame.columns if frame[column].dropna().empty
    ]
    if empty_columns:
        raise ProposalRejected("table-empty-column")
    return frame


def _is_bool_series(values: Any) -> bool:
    return bool(values) and all(isinstance(value, bool) for value in values)


def _column_profile(name: str, series: Any) -> ColumnProfile:
    nonempty = series.dropna()
    values = nonempty.tolist()
    explicit_temporal = bool(TEMPORAL_COLUMN_PATTERN.search(name))
    numeric_values = pd.to_numeric(nonempty, errors="coerce")
    numeric = (
        len(nonempty) > 0
        and not _is_bool_series(values)
        and numeric_values.notna().all()
    )
    temporal_valid = False
    if explicit_temporal and len(nonempty) > 0:
        if numeric:
            temporal_valid = True
        elif all(isinstance(value, (date, datetime)) for value in values):
            temporal_valid = True
        else:
            parsed = pd.to_datetime(
                nonempty.astype(str),
                errors="coerce",
                format="mixed",
            )
            temporal_valid = parsed.notna().all()
    monotonic = False
    if numeric and len(numeric_values) >= 2 and numeric_values.nunique() >= 2:
        monotonic = bool(
            numeric_values.is_monotonic_increasing
            or numeric_values.is_monotonic_decreasing
        )
    categorical = (
        len(nonempty) > 0
        and not numeric
        and not explicit_temporal
    )
    return ColumnProfile(
        name=name,
        numeric=numeric,
        explicit_temporal=explicit_temporal,
        temporal_valid=temporal_valid,
        categorical=categorical,
        monotonic_numeric=monotonic,
        non_null_count=len(nonempty),
        unique_count=int(nonempty.nunique(dropna=True)),
    )


def _looks_like_index_header(name: str) -> bool:
    normalized = str(name).strip().casefold()
    if not normalized:
        return False
    if normalized.endswith("#") or normalized in INDEX_HEADER_TOKENS:
        return True
    compact = re.sub(r"[._-]+", " ", normalized)
    parts = [part for part in compact.split() if part]
    if not parts or parts[0] not in INDEX_HEADER_WORDS:
        return False
    return (
        len(parts) == 1
        or parts[-1] in INDEX_HEADER_TOKENS
        or parts[-1] in {"no", "number", "id", "index"}
    )


def _is_replicate_index_column(name: str, series: Any) -> bool:
    """Recognize an integer replicate/row counter without guessing from values alone."""

    if not _looks_like_index_header(name):
        return False
    nonempty = series.dropna()
    if len(nonempty) < 2:
        return False
    numeric = pd.to_numeric(nonempty, errors="coerce")
    if numeric.isna().any() or not all(float(value).is_integer() for value in numeric):
        return False
    values = [int(value) for value in numeric]
    minimum = min(values)
    maximum = max(values)
    if minimum not in (0, 1) or set(values) != set(range(minimum, maximum + 1)):
        return False
    # A row counter may be globally shuffled while preserving the complete
    # 0/1-based range. It may also restart for each treatment block
    # (1,2,3,1,2,3), but cannot jump within a block.
    if len(set(values)) == len(values):
        return len(values) == maximum - minimum + 1
    return all(
        current == previous + 1 or current == minimum
        for previous, current in zip(values, values[1:], strict=False)
    )


def _is_identifier_column_v4(name: str, series: Any) -> bool:
    """Recognize a named integer identifier even when its order is arbitrary."""

    # A bare ``Count`` is conventionally a measured frequency, not an entity
    # identifier. Keep it eligible as y even when its values are small ints.
    if str(name).strip().casefold() == "count":
        return False
    if _is_replicate_index_column(name, series):
        return True
    if not _looks_like_index_header(name):
        return False
    nonempty = series.dropna()
    if len(nonempty) < 2:
        return False
    numeric = pd.to_numeric(nonempty, errors="coerce")
    if numeric.isna().any() or not all(float(value).is_integer() for value in numeric):
        return False
    values = [int(value) for value in numeric]
    if min(values) < 0:
        return False
    # Header evidence is required. Given it, a compact integer label domain is
    # an identifier even if source rows were sorted by a different variable.
    return max(values) <= max(50, len(values) * 4)


def _is_controlled_numeric_column(profile: ColumnProfile) -> bool:
    normalized = profile.name.strip().casefold()
    return normalized in {
        "msec",
        "ms",
        "sec",
        "second",
        "seconds",
        "min",
        "minute",
        "minutes",
        "hr",
        "hour",
        "hours",
    } or bool(
        CONTROLLED_NUMERIC_COLUMN_PATTERN.search(profile.name)
        or SCATTERING_VECTOR_COLUMN_PATTERN.search(profile.name)
    )


def _is_group_label_header(profile: ColumnProfile) -> bool:
    """Conservatively identify a wide-table column header as a group label."""

    name = profile.name.strip()
    normalized = name.casefold()
    if (
        not name
        or profile.explicit_temporal
        or _is_controlled_numeric_column(profile)
        or bool(MEASUREMENT_COLUMN_PATTERN.search(name))
    ):
        return False
    if any(marker in name for marker in ("+", ";", "/", "[", "]")):
        return True
    if normalized in {
        "wt",
        "ko",
        "control",
        "vehicle",
        "treated",
        "untreated",
        "mutant",
        "wildtype",
        "wild-type",
    }:
        return True
    compact = re.sub(r"[^a-z0-9]+", "", normalized)
    return bool(compact) and len(compact) <= 16 and not any(
        token in compact
        for token in (
            "value",
            "signal",
            "response",
            "measure",
            "intensity",
            "density",
            "ratio",
            "length",
            "distance",
            "volume",
            "weight",
            "height",
            "width",
        )
    )


def _is_v4_temporal_profile(profile: ColumnProfile) -> bool:
    if profile.temporal_valid:
        return True
    if not profile.numeric or not V4_TEMPORAL_NAME_PATTERN.search(profile.name):
        return False
    unit = _extract_unit(profile.name)
    return unit is None or unit.strip().casefold() in TEMPORAL_UNIT_NAMES


def _has_many_group_observations(
    frame: Any,
    *,
    x_column: str,
    y_columns: tuple[str, ...],
) -> bool:
    if len(y_columns) < 1:
        return False
    observed = frame[x_column].notna() & frame[list(y_columns)].notna().any(axis=1)
    counts = frame.loc[observed].groupby(x_column, dropna=True).size()
    return (
        len(counts) >= 2
        and not counts.empty
        and int(counts.min()) >= GROUPED_DOT_PLOT_MIN_VALUES
    )


def _v4_direct_group_supports_dot_plot(
    frame: Any,
    *,
    x_column: str,
    y_columns: tuple[str, ...],
    profiles: dict[str, ColumnProfile],
) -> bool:
    """Require a response-like y header for multi-series direct group tables."""

    if not _has_many_group_observations(
        frame,
        x_column=x_column,
        y_columns=y_columns,
    ):
        return False
    if len(y_columns) == 1:
        return True
    return any(
        bool(MEASUREMENT_COLUMN_PATTERN.search(name))
        or bool(DIRECT_GROUP_RESPONSE_PATTERN.search(name))
        or not _is_group_label_header(profiles[name])
        for name in y_columns
    )


def _v4_numeric_x_requires_scatter(
    frame: Any,
    profile: ColumnProfile,
) -> bool:
    if _is_controlled_numeric_column(profile):
        return False
    values = pd.to_numeric(frame[profile.name].dropna(), errors="coerce")
    return bool(values.notna().all() and values.duplicated().any())


def _extract_unit(column: str) -> str | None:
    for pattern in (PAREN_UNIT_PATTERN, BRACKET_UNIT_PATTERN):
        match = pattern.search(column)
        if match and match.group(1).strip():
            return match.group(1).strip()
    match = SUFFIX_UNIT_PATTERN.search(column)
    return match.group(1) if match else None


def _direct_table_analysis(
    *,
    headers: list[str],
    profiles: dict[str, ColumnProfile],
    rows: int,
    x_profile: ColumnProfile,
    x_mode: str,
    chart_family: str,
    excluded_index_columns: tuple[str, ...] = (),
    excluded_non_numeric_columns: tuple[str, ...] = (),
) -> TableAnalysis:
    excluded = set(excluded_index_columns) | set(excluded_non_numeric_columns)
    remaining = [
        profile
        for name, profile in profiles.items()
        if name != x_profile.name and name not in excluded
    ]
    if any(not profile.numeric for profile in remaining):
        raise ProposalRejected("non-numeric-y-column")
    if len(remaining) < 1:
        raise ProposalRejected("numeric-y-column-missing")
    if len(remaining) > 4:
        raise ProposalRejected("too-many-y-columns")
    units = {
        name: unit
        for name in headers
        if (unit := _extract_unit(name)) is not None
    }
    return TableAnalysis(
        x=x_profile.name,
        y=tuple(profile.name for profile in remaining),
        x_mode=x_mode,
        chart_family=chart_family,
        units=units,
        rows=rows,
        columns=len(headers),
        non_null_y={
            profile.name: profile.non_null_count for profile in remaining
        },
        dropped_index_columns=excluded_index_columns,
    )


def _analyze_table_v3_or_v4(
    frame: Any,
    *,
    headers: list[str],
    profiles: dict[str, ColumnProfile],
    v4: bool,
) -> TableAnalysis:
    invalid_temporal = [
        name
        for name, profile in profiles.items()
        if profile.explicit_temporal and not profile.temporal_valid
    ]
    if invalid_temporal:
        raise ProposalRejected("explicit-temporal-column-invalid")

    if v4:
        temporal = [
            profile
            for profile in profiles.values()
            if _is_v4_temporal_profile(profile)
        ]
    else:
        temporal = [
            profile for profile in profiles.values() if profile.temporal_valid
        ]
    categorical = [
        profile for profile in profiles.values() if profile.categorical
    ]
    run_order = [
        profile
        for profile in profiles.values()
        if RUN_ORDER_COLUMN_PATTERN.fullmatch(profile.name.strip())
    ]
    index_detector = _is_identifier_column_v4 if v4 else _is_replicate_index_column
    index_columns = tuple(
        name
        for name, profile in profiles.items()
        if profile.numeric and index_detector(name, frame[name])
    )

    if len(run_order) > 1:
        raise ProposalRejected("x-column-ambiguous-run-order")
    if run_order:
        x_profile = run_order[0]
        if not x_profile.monotonic_numeric:
            raise ProposalRejected("explicit-run-order-column-invalid")
        return _direct_table_analysis(
            headers=headers,
            profiles=profiles,
            rows=len(frame.index),
            x_profile=x_profile,
            x_mode="linear",
            chart_family="scatter",
            excluded_index_columns=tuple(
                name for name in index_columns if name != x_profile.name
            ),
        )
    if len(temporal) > 1:
        raise ProposalRejected("x-column-ambiguous-temporal")
    if temporal:
        x_profile = temporal[0]
        return _direct_table_analysis(
            headers=headers,
            profiles=profiles,
            rows=len(frame.index),
            x_profile=x_profile,
            x_mode="temporal",
            chart_family="line",
            excluded_index_columns=tuple(
                name for name in index_columns if name != x_profile.name
            ),
            excluded_non_numeric_columns=(
                tuple(
                    profile.name
                    for profile in profiles.values()
                    if profile.categorical
                )
                if v4
                else ()
            ),
        )
    if len(categorical) > 1:
        raise ProposalRejected("x-column-ambiguous-categorical")
    if len(categorical) == 1:
        x_profile = categorical[0]
        excluded_index_columns = index_columns
        y_columns = tuple(
            profile.name
            for name, profile in profiles.items()
            if name != x_profile.name
            and name not in excluded_index_columns
            and profile.numeric
        )
        return _direct_table_analysis(
            headers=headers,
            profiles=profiles,
            rows=len(frame.index),
            x_profile=x_profile,
            x_mode="categorical",
            chart_family=(
                "scatter"
                if v4
                and _v4_direct_group_supports_dot_plot(
                    frame,
                    x_column=x_profile.name,
                    y_columns=y_columns,
                    profiles=profiles,
                )
                else "bar"
            ),
            excluded_index_columns=excluded_index_columns,
        )

    measurements = [
        profile
        for name, profile in profiles.items()
        if profile.numeric and name not in index_columns
    ]
    if (
        len(measurements) >= 2
        and all(_is_group_label_header(profile) for profile in measurements)
    ):
        wide_columns = tuple(profile.name for profile in measurements)
        # V3 preserved a dot for any repeated group value. V4 distinguishes
        # aggregate tables from individual-point tables deterministically.
        if v4:
            chart_family = (
                "scatter"
                if all(
                    profile.non_null_count >= GROUPED_DOT_PLOT_MIN_VALUES
                    for profile in measurements
                )
                else "bar"
            )
        else:
            chart_family = (
                "bar"
                if all(profile.non_null_count <= 1 for profile in measurements)
                else "scatter"
            )
        return TableAnalysis(
            x=WIDE_MELT_GROUP_COLUMN,
            y=(WIDE_MELT_VALUE_COLUMN,),
            x_mode="categorical",
            chart_family=chart_family,
            units={},
            rows=len(frame.index),
            columns=len(headers),
            non_null_y={
                WIDE_MELT_VALUE_COLUMN: sum(
                    profile.non_null_count for profile in measurements
                )
            },
            binding_mode=WIDE_MELT_BINDING_MODE,
            wide_value_columns=wide_columns,
            dropped_index_columns=index_columns,
        )

    if len(measurements) == 2 and not any(
        _is_controlled_numeric_column(profile) for profile in measurements
    ):
        monotonic = [
            profile for profile in measurements if profile.monotonic_numeric
        ]
        x_profile = monotonic[0] if len(monotonic) == 1 else measurements[0]
        return _direct_table_analysis(
            headers=headers,
            profiles=profiles,
            rows=len(frame.index),
            x_profile=x_profile,
            x_mode="linear",
            chart_family="scatter",
            excluded_index_columns=index_columns,
        )

    controlled = [
        profile
        for profile in measurements
        if profile.monotonic_numeric
        and _is_controlled_numeric_column(profile)
    ]
    if len(controlled) == 1:
        return _direct_table_analysis(
            headers=headers,
            profiles=profiles,
            rows=len(frame.index),
            x_profile=controlled[0],
            x_mode="linear",
            chart_family="line",
            excluded_index_columns=index_columns,
        )

    # Preserve the v2 fallback for structures that are neither a confident
    # grouped-wide table nor a two-measurement correlation.
    monotonic = [
        profile
        for profile in profiles.values()
        if profile.monotonic_numeric
    ]
    if len(monotonic) != 1:
        raise ProposalRejected(
            "x-column-not-found"
            if not monotonic
            else "x-column-ambiguous-monotonic-numeric"
        )
    x_profile = monotonic[0]
    return _direct_table_analysis(
        headers=headers,
        profiles=profiles,
        rows=len(frame.index),
        x_profile=x_profile,
        x_mode="linear",
        chart_family=(
            "scatter"
            if v4 and _v4_numeric_x_requires_scatter(frame, x_profile)
            else "line"
        ),
        excluded_index_columns=tuple(
            name for name in index_columns if name != x_profile.name
        ),
    )


def _analyze_table_v3(
    frame: Any,
    *,
    headers: list[str],
    profiles: dict[str, ColumnProfile],
) -> TableAnalysis:
    return _analyze_table_v3_or_v4(
        frame,
        headers=headers,
        profiles=profiles,
        v4=False,
    )


def _analyze_table_v4(
    frame: Any,
    *,
    headers: list[str],
    profiles: dict[str, ColumnProfile],
) -> TableAnalysis:
    return _analyze_table_v3_or_v4(
        frame,
        headers=headers,
        profiles=profiles,
        v4=True,
    )


def analyze_table(
    frame: Any,
    *,
    rule_version: str = PROPOSAL_RULE_V1,
) -> TableAnalysis:
    if rule_version not in {
        PROPOSAL_RULE_V1,
        PROPOSAL_RULE_V2,
        PROPOSAL_RULE_V3,
        PROPOSAL_RULE_V4,
    }:
        raise ValueError(f"unsupported proposal rule version: {rule_version}")
    headers = _validate_headers(frame.columns)
    profiles = {
        name: _column_profile(name, frame[name])
        for name in headers
    }
    if rule_version == PROPOSAL_RULE_V3:
        return _analyze_table_v3(
            frame,
            headers=headers,
            profiles=profiles,
        )
    if rule_version == PROPOSAL_RULE_V4:
        return _analyze_table_v4(
            frame,
            headers=headers,
            profiles=profiles,
        )
    invalid_temporal = [
        name
        for name, profile in profiles.items()
        if profile.explicit_temporal and not profile.temporal_valid
    ]
    if invalid_temporal:
        raise ProposalRejected("explicit-temporal-column-invalid")

    temporal = [
        profile for profile in profiles.values() if profile.temporal_valid
    ]
    categorical = [
        profile for profile in profiles.values() if profile.categorical
    ]
    run_order = [
        profile
        for profile in profiles.values()
        if RUN_ORDER_COLUMN_PATTERN.fullmatch(profile.name.strip())
    ]
    if rule_version == PROPOSAL_RULE_V2 and len(run_order) > 1:
        raise ProposalRejected("x-column-ambiguous-run-order")
    if rule_version == PROPOSAL_RULE_V2 and run_order:
        x_profile = run_order[0]
        if not x_profile.monotonic_numeric:
            raise ProposalRejected("explicit-run-order-column-invalid")
        x_mode = "linear"
        chart_family = "scatter"
    elif len(temporal) > 1:
        raise ProposalRejected("x-column-ambiguous-temporal")
    elif temporal:
        x_profile = temporal[0]
        x_mode = "temporal"
        chart_family = "line"
    elif len(categorical) > 1:
        raise ProposalRejected("x-column-ambiguous-categorical")
    elif len(categorical) == 1:
        x_profile = categorical[0]
        x_mode = "categorical"
        chart_family = "bar"
    else:
        monotonic = [
            profile
            for profile in profiles.values()
            if profile.monotonic_numeric
        ]
        if len(monotonic) != 1:
            raise ProposalRejected(
                "x-column-not-found"
                if not monotonic
                else "x-column-ambiguous-monotonic-numeric"
            )
        x_profile = monotonic[0]
        x_mode = "linear"
        chart_family = "line"

    remaining = [
        profile
        for name, profile in profiles.items()
        if name != x_profile.name
    ]
    if any(not profile.numeric for profile in remaining):
        raise ProposalRejected("non-numeric-y-column")
    if len(remaining) < 1:
        raise ProposalRejected("numeric-y-column-missing")
    if len(remaining) > 4:
        raise ProposalRejected("too-many-y-columns")
    units = {
        name: unit
        for name in headers
        if (unit := _extract_unit(name)) is not None
    }
    return TableAnalysis(
        x=x_profile.name,
        y=tuple(profile.name for profile in remaining),
        x_mode=x_mode,
        chart_family=chart_family,
        units=units,
        rows=len(frame.index),
        columns=len(headers),
        non_null_y={
            profile.name: profile.non_null_count for profile in remaining
        },
    )


def _single_expectation(
    candidate: dict[str, Any],
    analysis: TableAnalysis,
) -> dict[str, Any]:
    panel_ids = candidate.get("panel_ids") or []
    if len(panel_ids) != 1:
        raise ProposalRejected("candidate-panel-id-not-unique")
    panel_id = str(panel_ids[0])
    series = []
    for y_column in analysis.y:
        descriptor: dict[str, Any] = {
            "series_id": y_column,
            "kind": analysis.chart_family,
            "x": analysis.x,
            "y": y_column,
        }
        series.append(descriptor)
    panel: dict[str, Any] = {
        "panel_id": panel_id,
        "axis_index": 0,
        "x_scale": "linear",
        "series": series,
    }
    if analysis.x in analysis.units:
        panel["x_unit"] = analysis.units[analysis.x]
    y_units = {
        analysis.units[y_column]
        for y_column in analysis.y
        if y_column in analysis.units
    }
    if (
        len(y_units) == 1
        and all(column in analysis.units for column in analysis.y)
    ):
        panel["y_unit"] = next(iter(y_units))
    return {
        "schema_version": EVALUATION_SCHEMA_VERSION,
        "panels": [panel],
        "panel_groups": [],
    }


def _user_goal(candidate: dict[str, Any], analysis: TableAnalysis) -> str:
    figure_no = candidate.get("figure_no")
    panel_id = str((candidate.get("panel_ids") or ["?"])[0]).upper()
    verb = {
        "line": "line chart",
        "bar": "bar chart",
        "scatter": "scatter plot",
    }[analysis.chart_family]
    if analysis.binding_mode == WIDE_MELT_BINDING_MODE:
        groups = ", ".join(f"`{name}`" for name in analysis.wide_value_columns)
        return (
            f"Create a {verb} for Figure {figure_no}{panel_id} showing values "
            f"grouped by the source-table column headers {groups}."
        )
    y_names = ", ".join(f"`{name}`" for name in analysis.y)
    return (
        f"Create a {verb} for Figure {figure_no}{panel_id} showing "
        f"{y_names} against `{analysis.x}`."
    )


def propose_single_candidate(
    candidate: dict[str, Any],
    *,
    input_candidates_sha256: str,
    code_commit: str,
    max_file_bytes: int,
    max_rows: int,
    max_columns: int,
    rule_version: str = PROPOSAL_RULE_V1,
) -> dict[str, Any]:
    frame = read_candidate_table(
        candidate,
        max_file_bytes=max_file_bytes,
        max_rows=max_rows,
        max_columns=max_columns,
    )
    analysis = analyze_table(frame, rule_version=rule_version)
    renderability = _renderability_audit(frame, analysis)
    if renderability["decision"] != "accepted":
        raise ProposalRejected(
            str(renderability["reason"]),
            detail=(
                "unique_categorical_x="
                f"{renderability['unique_categorical_x']};"
                f"max={MAX_BAR_CATEGORICAL_X};"
                f"policy={RENDERABILITY_POLICY_ID}"
            ),
            audit=renderability,
        )
    expectation = _single_expectation(candidate, analysis)
    source = candidate.get("source_table") or {}
    panel_id = str(candidate["panel_ids"][0])
    user_goal = _user_goal(candidate, analysis)
    intent = {
        "x": analysis.x,
        "y": analysis.y[0],
        "series": list(analysis.y),
        "x_scale": analysis.x_mode,
        "units": analysis.units,
    }
    if analysis.binding_mode == WIDE_MELT_BINDING_MODE:
        intent.update(
            {
                "binding_mode": WIDE_MELT_BINDING_MODE,
                "wide_melt": {
                    "group_column": WIDE_MELT_GROUP_COLUMN,
                    "value_column": WIDE_MELT_VALUE_COLUMN,
                    "source_value_columns": list(analysis.wide_value_columns),
                    "dropped_index_columns": list(
                        analysis.dropped_index_columns
                    ),
                },
            }
        )
    experiment_case = {
        "case_id": candidate.get("candidate_id"),
        "panel_count": 1,
        "split": None,
        "data_path": source.get("path"),
        "sheet": source.get("sheet_name"),
        "panel_id": panel_id,
        "panels": [
            {
                "panel_id": panel_id,
                "data_path": source.get("path"),
                "sheet": source.get("sheet_name"),
                "user_goal": user_goal,
            }
        ],
        "user_goal": user_goal,
        "chart_family": analysis.chart_family,
        "intent": intent,
        "evaluation_expectation": expectation,
    }
    proposal = deepcopy(candidate)
    proposal.update(
        {
            "proposal_type": "single_panel",
            "proposal_rule_version": rule_version,
            "renderability_policy_id": RENDERABILITY_POLICY_ID,
            "renderability_policy_hash": RENDERABILITY_POLICY_HASH,
            "renderability_audit": renderability,
            "curation_status": "proposed",
            "eligible_for_experiment": False,
            "eligibility_reasons": ["external-validation-required"],
            "experiment_case": experiment_case,
            "proposal_analysis": {
                "rows": analysis.rows,
                "columns": analysis.columns,
                "x": analysis.x,
                "y": list(analysis.y),
                "x_mode": analysis.x_mode,
                "chart_family": analysis.chart_family,
                "unique_categorical_x": renderability[
                    "unique_categorical_x"
                ],
                "units": analysis.units,
                "non_null_y": analysis.non_null_y,
                "binding_mode": analysis.binding_mode,
                "wide_value_columns": list(analysis.wide_value_columns),
                "dropped_index_columns": list(
                    analysis.dropped_index_columns
                ),
            },
            "input_candidates_sha256": input_candidates_sha256,
            "code_commit": code_commit,
        }
    )
    return proposal


def _validated_proposal_renderability(
    proposal: dict[str, Any],
) -> dict[str, Any]:
    frame = read_candidate_table(proposal)
    analysis = analyze_table(
        frame,
        rule_version=_resolve_proposal_rule_version(proposal),
    )
    computed = _renderability_audit(frame, analysis)
    declared = proposal.get("renderability_audit")
    if declared is not None and declared != computed:
        raise ValueError("proposal-renderability-audit-mismatch")
    if computed["decision"] != "accepted":
        raise ProposalRejected(
            str(computed["reason"]),
            detail=(
                "multi-constituent-unique_categorical_x="
                f"{computed['unique_categorical_x']};"
                f"max={MAX_BAR_CATEGORICAL_X}"
            ),
            audit=computed,
        )
    return computed


def _multi_panel_proposals(
    singles: list[dict[str, Any]],
    *,
    input_candidates_sha256: str,
    code_commit: str,
    blocked_groups: set[tuple[str, int]] | None = None,
) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, int], list[dict[str, Any]]] = {}
    rule_versions: dict[int, str] = {}
    renderability_audits: dict[int, dict[str, Any]] = {}
    for proposal in singles:
        rule_versions[id(proposal)] = _resolve_proposal_rule_version(proposal)
        renderability_audits[id(proposal)] = (
            _validated_proposal_renderability(proposal)
        )
        doi = str(proposal.get("doi") or "")
        figure_no = proposal.get("figure_no")
        if doi and isinstance(figure_no, int):
            grouped.setdefault((doi, figure_no), []).append(proposal)

    results: list[dict[str, Any]] = []
    for (doi, figure_no), group in sorted(grouped.items()):
        if blocked_groups and (doi, figure_no) in blocked_groups:
            continue
        by_panel: dict[str, list[dict[str, Any]]] = {}
        for proposal in group:
            panel_id = str(proposal["panel_ids"][0])
            by_panel.setdefault(panel_id, []).append(proposal)
        clear = [
            proposals[0]
            for panel_id, proposals in sorted(by_panel.items())
            if len(proposals) == 1
        ]
        if len(clear) < 2:
            continue
        group_rule_versions = {
            rule_versions[id(proposal)] for proposal in clear
        }
        if len(group_rule_versions) != 1:
            raise ValueError("multi-panel-proposal-rule-version-mixed")
        rule_version = next(iter(group_rule_versions))
        source_ids = sorted(str(item["candidate_id"]) for item in clear)
        digest = hashlib.sha256(
            "|".join(source_ids).encode("utf-8")
        ).hexdigest()[:12]
        article_id = doi.split("/", 1)[-1]
        case_id = f"multi-{article_id}-figure{figure_no}-{digest}"
        panels = []
        panel_renderability = []
        expectation_panels = []
        x_scales: list[str] = []
        x_units: list[str | None] = []
        series_sets: list[set[str]] = []
        for proposal in clear:
            case = proposal["experiment_case"]
            panel_id = str(proposal["panel_ids"][0])
            panels.append(
                {
                    "id": panel_id,
                    "data_path": case["data_path"],
                    "sheet": case.get("sheet"),
                    "user_goal": case["user_goal"],
                    "chart_family": case["chart_family"],
                    "intent": case["intent"],
                }
            )
            panel_renderability.append(
                {
                    "panel_id": panel_id,
                    **renderability_audits[id(proposal)],
                }
            )
            panel_expectation = case["evaluation_expectation"]["panels"][0]
            expectation_panels.append(panel_expectation)
            x_scales.append(str(panel_expectation.get("x_scale") or ""))
            x_units.append(panel_expectation.get("x_unit"))
            series_sets.append(
                {str(series["series_id"]) for series in panel_expectation["series"]}
            )

        checks: dict[str, Any] = {}
        cohesion: dict[str, Any] = {
            "group_id": "all_panels",
            "panels": [panel["id"] for panel in panels],
            "checks": checks,
        }
        if len(set(x_scales)) == 1:
            checks["shared_x_scale"] = True
        if all(unit is not None for unit in x_units) and len(set(x_units)) == 1:
            checks["shared_x_unit"] = True
        shared_series = sorted(set.intersection(*series_sets)) if series_sets else []
        if shared_series:
            cohesion["series"] = shared_series
            checks["palette_consistent"] = True
        panel_groups = [cohesion] if checks else []
        expectation = {
            "schema_version": EVALUATION_SCHEMA_VERSION,
            "panels": expectation_panels,
            "panel_groups": panel_groups,
        }
        experiment_case = {
            "case_id": case_id,
            "panel_count": len(panels),
            "split": None,
            "panels": panels,
            "user_goal": (
                f"Create a {len(panels)}-panel Figure {figure_no} using "
                "the proposed panel-specific charts."
            ),
            "chart_family": "multi_panel",
            "intent": {
                "panels": [
                    {
                        "panel_id": panel["id"],
                        **panel["intent"],
                    }
                    for panel in panels
                ]
            },
            "evaluation_expectation": expectation,
        }
        results.append(
            {
                "schema_version": SCHEMA_VERSION,
                "candidate_id": case_id,
                "proposal_type": "multi_panel",
                "proposal_rule_version": rule_version,
                "renderability_policy_id": RENDERABILITY_POLICY_ID,
                "renderability_policy_hash": RENDERABILITY_POLICY_HASH,
                "renderability_audit": {
                    "schema_version": SCHEMA_VERSION,
                    "policy_id": RENDERABILITY_POLICY_ID,
                    "policy_hash": RENDERABILITY_POLICY_HASH,
                    "decision": "accepted",
                    "reason": "all_constituent_panels_pass",
                    "multi_panel_rule": "reject_case_if_any_panel_fails",
                    "panels": panel_renderability,
                },
                "source_candidate_ids": source_ids,
                "doi": doi,
                "figure_no": figure_no,
                "panel_ids": [panel["id"] for panel in panels],
                "curation_status": "proposed",
                "eligible_for_experiment": False,
                "eligibility_reasons": ["external-validation-required"],
                "experiment_case": experiment_case,
                "input_candidates_sha256": input_candidates_sha256,
                "code_commit": code_commit,
            }
        )
    return sorted(results, key=lambda item: item["candidate_id"])


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        value = json.loads(line)
        if not isinstance(value, dict):
            raise ValueError(f"{path}:{line_number}: expected an object")
        records.append(value)
    return records


def resolve_code_commit(repository: str | Path | None = None) -> str:
    cwd = Path(repository) if repository else Path(__file__).resolve().parents[2]
    try:
        result = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=cwd,
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return "unknown"
    return result.stdout.strip() or "unknown"


def propose_cases(
    *,
    candidates_path: str | Path,
    code_commit: str | None = None,
    max_file_bytes: int = DEFAULT_MAX_FILE_BYTES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
    rule_version: str = CURRENT_PROPOSAL_RULE,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    if min(max_file_bytes, max_rows, max_columns) < 1:
        raise ValueError("proposal limits must be positive")
    if rule_version not in {
        PROPOSAL_RULE_V1,
        PROPOSAL_RULE_V2,
        PROPOSAL_RULE_V3,
        PROPOSAL_RULE_V4,
    }:
        raise ValueError(f"unsupported proposal rule version: {rule_version}")
    _validate_renderability_policy()
    source_path = Path(candidates_path)
    input_hash = sha256_file(source_path)
    commit = code_commit or resolve_code_commit()
    candidates = sorted(
        _read_jsonl(source_path),
        key=lambda item: str(item.get("candidate_id") or ""),
    )
    singles: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    blocked_multi_groups: set[tuple[str, int]] = set()
    renderability_rejected = 0
    verified_preserved = 0
    for candidate in candidates:
        status = str(candidate.get("curation_status") or "").casefold()
        if status == "verified":
            verified_preserved += 1
            record = deepcopy(candidate)
            record.update(
                {
                    "proposal_rejection_reasons": [
                        "verified-candidate-preserved"
                    ],
                    "input_candidates_sha256": input_hash,
                    "code_commit": commit,
                }
            )
            rejected.append(record)
            continue
        if status != "unverified":
            record = deepcopy(candidate)
            record.update(
                {
                    "proposal_rejection_reasons": [
                        "candidate-not-unverified"
                    ],
                    "input_candidates_sha256": input_hash,
                    "code_commit": commit,
                }
            )
            rejected.append(record)
            continue
        try:
            singles.append(
                propose_single_candidate(
                    candidate,
                    input_candidates_sha256=input_hash,
                    code_commit=commit,
                    max_file_bytes=max_file_bytes,
                    max_rows=max_rows,
                    max_columns=max_columns,
                    rule_version=rule_version,
                )
            )
        except ProposalRejected as exc:
            record = deepcopy(candidate)
            rejection_fields: dict[str, Any] = {
                "proposal_rejection_reasons": list(exc.reasons),
                "proposal_rejection_detail": exc.detail,
                "input_candidates_sha256": input_hash,
                "code_commit": commit,
            }
            if exc.audit is not None:
                rejection_fields.update(
                    {
                        "renderability_policy_id": RENDERABILITY_POLICY_ID,
                        "renderability_policy_hash": RENDERABILITY_POLICY_HASH,
                        "renderability_audit": exc.audit,
                    }
                )
                renderability_rejected += 1
                group_key = _candidate_group_key(candidate)
                if group_key is not None:
                    blocked_multi_groups.add(group_key)
            record.update(rejection_fields)
            rejected.append(record)
    singles.sort(key=lambda item: item["candidate_id"])
    multi = _multi_panel_proposals(
        singles,
        input_candidates_sha256=input_hash,
        code_commit=commit,
        blocked_groups=blocked_multi_groups,
    )
    proposed = sorted(
        singles + multi,
        key=lambda item: (item["proposal_type"], item["candidate_id"]),
    )
    rejected.sort(key=lambda item: str(item.get("candidate_id") or ""))
    if any(item.get("eligible_for_experiment") for item in proposed):
        raise AssertionError("proposals must never be experiment-eligible")
    for item in proposed:
        audit = item.get("renderability_audit")
        if not isinstance(audit, dict) or audit.get("decision") != "accepted":
            raise AssertionError("proposed case has no passing renderability audit")
        panel_audits = (
            audit.get("panels")
            if item.get("proposal_type") == "multi_panel"
            else [audit]
        )
        if not isinstance(panel_audits, list):
            raise AssertionError("proposal renderability panels are invalid")
        for panel_audit in panel_audits:
            if not isinstance(panel_audit, dict):
                raise AssertionError("proposal renderability audit is invalid")
            count = panel_audit.get("unique_categorical_x")
            if isinstance(count, int) and count > MAX_BAR_CATEGORICAL_X:
                raise AssertionError("proposal crosses categorical bar cap")
    summary = {
        "schema_version": SCHEMA_VERSION,
        "input_candidates": str(source_path.resolve()),
        "input_candidates_sha256": input_hash,
        "code_commit": commit,
        "proposal_rule_version": rule_version,
        "renderability_policy_id": RENDERABILITY_POLICY_ID,
        "renderability_policy_hash": RENDERABILITY_POLICY_HASH,
        "max_unique_categorical_x_per_bar_panel": (
            MAX_BAR_CATEGORICAL_X
        ),
        "input_count": len(candidates),
        "single_proposals": len(singles),
        "multi_panel_proposals": len(multi),
        "proposals_total": len(proposed),
        "rejected": len(rejected),
        "renderability_rejected": renderability_rejected,
        "multi_groups_blocked_by_renderability": len(
            blocked_multi_groups
        ),
        "verified_preserved": verified_preserved,
        "eligible_for_experiment": 0,
        "max_file_bytes": max_file_bytes,
        "max_rows": max_rows,
        "max_columns": max_columns,
        "llm_calls": 0,
    }
    return proposed, rejected, summary


def _write_jsonl(path: Path, records: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(
                json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n"
            )


def write_proposal_outputs(
    output_root: str | Path,
    proposed: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    output = Path(output_root)
    output.mkdir(parents=True, exist_ok=True)
    _write_jsonl(output / "proposed.jsonl", proposed)
    _write_jsonl(output / "rejected.jsonl", rejected)
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
