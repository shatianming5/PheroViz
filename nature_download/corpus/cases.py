"""Build fail-closed benchmark candidates from licensed Source Data assets."""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
import json
from pathlib import Path, PurePosixPath
import re
import shutil
import stat
from typing import Any, Iterable
import zipfile

try:
    from openpyxl import load_workbook
except ImportError:  # surfaced as an explicit workbook inspection error
    load_workbook = None

from .policy import normalize_cc_by_url, normalize_doi
from .provenance import sha256_bytes, sha256_file


SCHEMA_VERSION = "1.0"
TABLE_SUFFIXES = frozenset({".csv", ".xlsx"})
FIGURE_SUFFIXES = frozenset(
    {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}
)
DEFAULT_MAX_ZIP_FILES = 1_000
DEFAULT_MAX_ZIP_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
DEFAULT_MAX_XLSX_SHEETS = 256
FIGURE_PANEL_PATTERN = re.compile(
    r"(?i)(?<![a-z0-9])(?:figure|fig)[\s_.-]*(?P<figure>\d{1,3})"
    r"[\s_.-]*(?P<panel>[a-z])(?=$|[\s_.-])"
)
MULTI_PANEL_PATTERN = re.compile(
    r"(?ix)"
    r"(?<![a-z0-9])(?:figure|fig)[\s_.-]*\d{1,3}[\s_.-]*"
    r"[a-z](?:"
    r"\s*(?:-|–|—|,|/|&|\band\b|\bto\b)\s*[a-z](?=$|[\s_.-])"
    r"|[a-z](?=$|[\s_.-])"
    r")"
)
SUPPLEMENTARY_PATTERN = re.compile(
    r"(?i)(?:supp(?:lementary)?|extended[\s_.-]*data)"
)
EVIDENCE_TYPES = frozenset({"human_review", "external_validation"})


class ArchiveSafetyError(ValueError):
    """Raised when an archive violates a fail-closed extraction rule."""

    def __init__(self, code: str, detail: str | None = None) -> None:
        self.code = code
        self.detail = detail
        message = code if not detail else f"{code}:{detail}"
        super().__init__(message)


class WorkbookInspectionError(ValueError):
    """Raised when an XLSX cannot be inspected without loading cells."""

    def __init__(self, code: str, detail: str | None = None) -> None:
        self.code = code
        self.detail = detail
        message = code if not detail else f"{code}:{detail}"
        super().__init__(message)


@dataclass(frozen=True)
class ExtractedTable:
    path: Path
    member_name: str


@dataclass(frozen=True)
class FilenameMapping:
    figure_no: int
    panel_id: str
    qualifier: str | None


@dataclass(frozen=True)
class WorkbookSheets:
    names: tuple[str, ...]
    total_count: int
    resource_count: int


@dataclass(frozen=True)
class MappingAttempt:
    mapping: FilenameMapping | None
    reasons: tuple[str, ...]
    sheet_name: str | None
    detail: str | None = None


def _canonical_json_sha256(value: Any) -> str:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return sha256_bytes(payload)


def _is_resource_fork_name(name: str) -> bool:
    normalized = str(name).replace("\\", "/")
    parts = [part for part in normalized.split("/") if part]
    return any(part.casefold() == "__macosx" for part in parts) or (
        bool(parts) and parts[-1].startswith("._")
    )


def _safe_member_path(name: str) -> PurePosixPath:
    if not name or "\x00" in name:
        raise ArchiveSafetyError("zip-invalid-member-name", repr(name))
    normalized = name.replace("\\", "/")
    if (
        normalized.startswith("/")
        or normalized.startswith("//")
        or re.match(r"^[a-zA-Z]:", normalized)
    ):
        raise ArchiveSafetyError("zip-absolute-path", name)
    path = PurePosixPath(normalized)
    if path.is_absolute() or any(part == ".." for part in path.parts):
        raise ArchiveSafetyError("zip-slip", name)
    if not path.parts or all(part in {"", "."} for part in path.parts):
        raise ArchiveSafetyError("zip-invalid-member-name", name)
    return path


def _member_mode(info: zipfile.ZipInfo) -> int:
    return (info.external_attr >> 16) & 0xFFFF


def _validate_zip_members(
    infos: list[zipfile.ZipInfo],
    *,
    max_files: int,
    max_uncompressed_bytes: int,
) -> list[tuple[zipfile.ZipInfo, PurePosixPath]]:
    if len(infos) > max_files:
        raise ArchiveSafetyError(
            "zip-file-count-limit",
            f"{len(infos)}>{max_files}",
        )
    files: list[tuple[zipfile.ZipInfo, PurePosixPath]] = []
    total_size = 0
    seen: set[str] = set()
    for info in infos:
        path = _safe_member_path(info.filename)
        mode = _member_mode(info)
        if stat.S_ISLNK(mode):
            raise ArchiveSafetyError("zip-symlink", info.filename)
        file_type = stat.S_IFMT(mode)
        if file_type not in {0, stat.S_IFREG, stat.S_IFDIR}:
            raise ArchiveSafetyError("zip-special-file", info.filename)
        if info.flag_bits & 0x1:
            raise ArchiveSafetyError("zip-encrypted-member", info.filename)
        if info.is_dir():
            continue
        key = path.as_posix().casefold()
        if key in seen:
            raise ArchiveSafetyError("zip-duplicate-member", info.filename)
        seen.add(key)
        if info.file_size < 0:
            raise ArchiveSafetyError("zip-invalid-member-size", info.filename)
        total_size += info.file_size
        files.append((info, path))
        if len(files) > max_files:
            raise ArchiveSafetyError(
                "zip-file-count-limit",
                f"{len(files)}>{max_files}",
            )
        if total_size > max_uncompressed_bytes:
            raise ArchiveSafetyError(
                "zip-uncompressed-size-limit",
                f"{total_size}>{max_uncompressed_bytes}",
            )
    return files


def safe_extract_tables(
    archive_path: str | Path,
    destination: str | Path,
    *,
    max_files: int = DEFAULT_MAX_ZIP_FILES,
    max_uncompressed_bytes: int = DEFAULT_MAX_ZIP_UNCOMPRESSED_BYTES,
) -> list[ExtractedTable]:
    """Extract only CSV/XLSX members after validating the entire archive."""
    if max_files < 1 or max_uncompressed_bytes < 1:
        raise ValueError("ZIP limits must be positive")
    archive = Path(archive_path)
    destination_path = Path(destination)
    if archive.is_symlink():
        raise ArchiveSafetyError("zip-input-symlink", str(archive))
    if not archive.is_file():
        raise ArchiveSafetyError("zip-input-missing", str(archive))

    try:
        with zipfile.ZipFile(archive) as handle:
            members = _validate_zip_members(
                handle.infolist(),
                max_files=max_files,
                max_uncompressed_bytes=max_uncompressed_bytes,
            )
            shutil.rmtree(destination_path, ignore_errors=True)
            destination_path.mkdir(parents=True, exist_ok=True)
            extracted: list[ExtractedTable] = []
            actual_total = 0
            for info, member_path in members:
                if _is_resource_fork_name(member_path.as_posix()):
                    continue
                if Path(member_path.name).suffix.casefold() not in TABLE_SUFFIXES:
                    continue
                target = destination_path.joinpath(*member_path.parts)
                target.parent.mkdir(parents=True, exist_ok=True)
                target_resolved = target.resolve()
                try:
                    target_resolved.relative_to(destination_path.resolve())
                except ValueError:
                    raise ArchiveSafetyError("zip-slip", info.filename) from None
                written = 0
                with handle.open(info, "r") as source, target.open("wb") as output:
                    while True:
                        chunk = source.read(1024 * 1024)
                        if not chunk:
                            break
                        written += len(chunk)
                        actual_total += len(chunk)
                        if (
                            written > info.file_size
                            or actual_total > max_uncompressed_bytes
                        ):
                            raise ArchiveSafetyError(
                                "zip-uncompressed-size-limit",
                                info.filename,
                            )
                        output.write(chunk)
                if written != info.file_size:
                    raise ArchiveSafetyError(
                        "zip-member-size-mismatch",
                        info.filename,
                    )
                extracted.append(
                    ExtractedTable(
                        path=target,
                        member_name=member_path.as_posix(),
                    )
                )
            return sorted(extracted, key=lambda item: item.member_name.casefold())
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        shutil.rmtree(destination_path, ignore_errors=True)
        if isinstance(exc, ArchiveSafetyError):
            raise
        raise ArchiveSafetyError("zip-invalid", str(exc)) from exc
    except ArchiveSafetyError:
        shutil.rmtree(destination_path, ignore_errors=True)
        raise


def parse_figure_panel(filename: str) -> tuple[FilenameMapping | None, list[str]]:
    basename = str(filename).replace("\\", "/").rsplit("/", 1)[-1]
    suffix = Path(basename).suffix.casefold()
    stem = basename[: -len(suffix)] if suffix in TABLE_SUFFIXES else basename
    if SUPPLEMENTARY_PATTERN.search(stem):
        return None, ["supplementary-figure-mapping-unsupported"]
    if MULTI_PANEL_PATTERN.search(stem):
        return None, ["multiple-panel-reference"]
    matches = list(FIGURE_PANEL_PATTERN.finditer(stem))
    unique = {
        (int(match.group("figure")), match.group("panel").casefold())
        for match in matches
    }
    if not unique:
        return None, ["figure-panel-pattern-not-found"]
    if len(unique) != 1:
        return None, ["multiple-conflicting-figure-panel-matches"]
    figure_no, panel_id = next(iter(unique))
    if figure_no < 1:
        return None, ["figure-number-invalid"]
    match = matches[0]
    qualifier = stem[match.end() :].strip(" ._-") or None
    return FilenameMapping(figure_no, panel_id, qualifier), []


def _read_xlsx_sheet_names(
    path: Path,
    *,
    max_sheets: int,
) -> WorkbookSheets:
    if max_sheets < 1:
        raise ValueError("max_sheets must be positive")
    if load_workbook is None:
        raise WorkbookInspectionError(
            "xlsx-openpyxl-missing",
            "install openpyxl from nature_download/requirements.txt",
        )
    workbook = None
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        sheet_names = tuple(str(name) for name in workbook.sheetnames)
    except Exception as exc:
        raise WorkbookInspectionError(
            "xlsx-workbook-invalid",
            f"{type(exc).__name__}:{exc}",
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
    if len(sheet_names) > max_sheets:
        raise WorkbookInspectionError(
            "xlsx-sheet-count-limit",
            f"{len(sheet_names)}>{max_sheets}",
        )
    usable = tuple(
        name for name in sheet_names if not _is_resource_fork_name(name)
    )
    return WorkbookSheets(
        names=usable,
        total_count=len(sheet_names),
        resource_count=len(sheet_names) - len(usable),
    )


def _mapping_attempts(
    path: Path,
    *,
    mapping_name: str,
    max_xlsx_sheets: int,
) -> tuple[list[MappingAttempt], dict[str, int]]:
    filename_mapping, filename_reasons = parse_figure_panel(mapping_name)
    metrics = {
        "xlsx_workbooks": 0,
        "xlsx_sheets_inspected": 0,
        "xlsx_resource_sheets_skipped": 0,
        "xlsx_corrupt": 0,
        "xlsx_sheet_limit_exceeded": 0,
    }
    if path.suffix.casefold() != ".xlsx":
        return [
            MappingAttempt(
                mapping=filename_mapping,
                reasons=tuple(filename_reasons),
                sheet_name=None,
            )
        ], metrics

    metrics["xlsx_workbooks"] = 1
    try:
        sheets = _read_xlsx_sheet_names(path, max_sheets=max_xlsx_sheets)
    except WorkbookInspectionError as exc:
        if exc.code == "xlsx-sheet-count-limit":
            metrics["xlsx_sheet_limit_exceeded"] = 1
        elif exc.code == "xlsx-workbook-invalid":
            metrics["xlsx_corrupt"] = 1
        return [
            MappingAttempt(
                mapping=filename_mapping,
                reasons=(exc.code,),
                sheet_name=None,
                detail=exc.detail,
            )
        ], metrics

    metrics["xlsx_sheets_inspected"] = sheets.total_count
    metrics["xlsx_resource_sheets_skipped"] = sheets.resource_count
    if not sheets.names:
        return [
            MappingAttempt(
                mapping=None,
                reasons=("xlsx-no-usable-sheets",),
                sheet_name=None,
            )
        ], metrics
    if filename_mapping is not None and not filename_reasons:
        return [
            MappingAttempt(
                mapping=filename_mapping,
                reasons=(),
                sheet_name=None,
            )
        ], metrics
    if "supplementary-figure-mapping-unsupported" in filename_reasons:
        return [
            MappingAttempt(
                mapping=None,
                reasons=tuple(filename_reasons),
                sheet_name=None,
            )
        ], metrics
    attempts: list[MappingAttempt] = []
    for sheet_name in sheets.names:
        mapping, reasons = parse_figure_panel(sheet_name)
        attempts.append(
            MappingAttempt(
                mapping=mapping,
                reasons=tuple(reasons),
                sheet_name=sheet_name,
            )
        )
    return attempts, metrics


def _read_manifest(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8")
    stripped = text.lstrip()
    if not stripped:
        return []
    if stripped.startswith("["):
        value = json.loads(text)
        if not isinstance(value, list):
            raise ValueError("corpus manifest JSON must be a list")
        return [item for item in value if isinstance(item, dict)]
    if stripped.startswith("{"):
        try:
            value = json.loads(text)
        except json.JSONDecodeError:
            value = None
        if isinstance(value, dict):
            for key in ("articles", "records", "manifests"):
                records = value.get(key)
                if isinstance(records, list):
                    return [item for item in records if isinstance(item, dict)]
            return [value]
    records: list[dict[str, Any]] = []
    for line_number, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        value = json.loads(line)
        if not isinstance(value, dict):
            raise ValueError(
                f"{path}:{line_number}: manifest row must be an object"
            )
        records.append(value)
    return records


def _article_id(record: dict[str, Any]) -> str | None:
    article_url = str(record.get("article_url") or "")
    match = re.search(r"/articles/([^/?#]+)", article_url)
    if match:
        return match.group(1)
    doi = normalize_doi(record.get("doi"))
    return doi.split("/", 1)[1] if doi and "/" in doi else None


def _load_evidence(path: str | Path | None) -> tuple[dict[str, dict[str, Any]], int]:
    if not path:
        return {}, 0
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(value, dict) and isinstance(value.get("verifications"), list):
        records = value["verifications"]
    elif isinstance(value, list):
        records = value
    elif isinstance(value, dict):
        records = []
        for candidate_id, evidence in value.items():
            if isinstance(evidence, dict):
                records.append({"candidate_id": candidate_id, **evidence})
    else:
        raise ValueError("evidence file must contain a list or object")

    accepted: dict[str, dict[str, Any]] = {}
    invalid = 0
    for record in records:
        if not isinstance(record, dict):
            invalid += 1
            continue
        candidate_id = str(record.get("candidate_id") or "").strip()
        status = str(
            record.get("curation_status") or record.get("status") or ""
        ).casefold()
        evidence_type = str(record.get("evidence_type") or "").casefold()
        evidence_ref = str(record.get("evidence_ref") or "").strip()
        reviewer_or_source = str(
            record.get("reviewer_or_source") or ""
        ).strip()
        experiment_case = record.get("experiment_case")
        case_complete = (
            isinstance(experiment_case, dict)
            and isinstance(experiment_case.get("user_goal"), str)
            and bool(experiment_case["user_goal"].strip())
            and isinstance(experiment_case.get("chart_family"), str)
            and bool(experiment_case["chart_family"].strip())
            and isinstance(
                experiment_case.get("evaluation_expectation"),
                dict,
            )
            and experiment_case["evaluation_expectation"].get(
                "schema_version"
            )
            == "1.1.0"
            and isinstance(
                experiment_case["evaluation_expectation"].get("panels"),
                list,
            )
            and bool(experiment_case["evaluation_expectation"]["panels"])
        )
        if (
            not candidate_id
            or status != "verified"
            or evidence_type not in EVIDENCE_TYPES
            or not evidence_ref
            or not reviewer_or_source
            or not case_complete
        ):
            invalid += 1
            continue
        accepted[candidate_id] = dict(record)
    return accepted, invalid


def _article_provenance(
    path: Path,
    *,
    doi: str,
    license_url: str,
) -> tuple[str | None, list[str]]:
    if not path.is_file():
        return None, ["article-provenance-manifest-missing"]
    digest = sha256_file(path)
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return digest, ["article-provenance-manifest-invalid"]
    if not isinstance(value, dict):
        return digest, ["article-provenance-manifest-invalid"]
    reasons: list[str] = []
    if normalize_doi(value.get("doi")) != doi:
        reasons.append("article-provenance-doi-mismatch")
    if not value.get("download_eligible"):
        reasons.append("article-provenance-not-download-eligible")
    provenance_license = value.get("license") or {}
    provenance_url = provenance_license.get("normalized_url") or provenance_license.get(
        "url"
    )
    if normalize_cc_by_url(provenance_url) != normalize_cc_by_url(license_url):
        reasons.append("article-provenance-license-mismatch")
    return digest, reasons


def _file_descriptor(path: Path, *, source_url: str | None = None) -> dict[str, Any]:
    return {
        "path": str(path.resolve()),
        "sha256": sha256_file(path),
        "size_bytes": path.stat().st_size,
        "source_url": source_url,
    }


def _figure_assets(
    article_dir: Path,
    figure_no: int,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None, list[str]]:
    figures_dir = article_dir / "figures"
    prefix = f"fig_{figure_no:03d}"
    image_paths = sorted(
        path
        for path in figures_dir.glob(f"{prefix}.*")
        if path.is_file() and path.suffix.casefold() in FIGURE_SUFFIXES
    )
    caption_paths = sorted(
        path
        for path in figures_dir.glob(f"{prefix}.txt")
        if path.is_file()
    )
    reasons: list[str] = []
    if not image_paths:
        reasons.append("figure-file-missing")
    elif len(image_paths) > 1:
        reasons.append("multiple-figure-files")
    if not caption_paths:
        reasons.append("caption-file-missing")
    elif len(caption_paths) > 1:
        reasons.append("multiple-caption-files")
    figure = _file_descriptor(image_paths[0]) if len(image_paths) == 1 else None
    caption = (
        _file_descriptor(caption_paths[0]) if len(caption_paths) == 1 else None
    )
    return figure, caption, reasons


def _source_table_descriptor(
    path: Path,
    *,
    content_root: Path,
    output_root: Path,
    archive_path: Path | None,
    archive_member: str | None,
    sheet_name: str | None,
) -> dict[str, Any]:
    root_name = "output_root" if archive_path else "content_root"
    root = output_root if archive_path else content_root
    descriptor = _file_descriptor(path)
    descriptor.update(
        {
            "format": path.suffix.casefold().lstrip("."),
            "path_root": root_name,
            "relative_path": path.resolve().relative_to(root.resolve()).as_posix(),
            "sheet_name": sheet_name,
            "archive": None,
        }
    )
    if archive_path:
        descriptor["archive"] = {
            **_file_descriptor(archive_path),
            "member_path": archive_member,
        }
    return descriptor


def _candidate_id(
    *,
    article_id: str,
    doi: str,
    mapping: FilenameMapping,
    table: dict[str, Any],
) -> str:
    identity = "|".join(
        [
            doi,
            str(mapping.figure_no),
            mapping.panel_id,
            mapping.qualifier or "",
            str(table["relative_path"]),
            str(table.get("sheet_name") or ""),
            str(table["sha256"]),
        ]
    )
    suffix = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:12]
    qualifier = (
        "-" + re.sub(r"[^a-z0-9]+", "-", mapping.qualifier.casefold()).strip("-")
        if mapping.qualifier
        else ""
    )
    return (
        f"case-{article_id}-figure{mapping.figure_no}-"
        f"panel{mapping.panel_id}{qualifier}-{suffix}"
    )


def _case_skeleton(
    *,
    candidate_id: str,
    mapping: FilenameMapping,
    table: dict[str, Any],
    curation_status: str,
    verification: dict[str, Any] | None,
) -> dict[str, Any]:
    eligible = curation_status == "verified"
    curated = (
        dict(verification["experiment_case"])
        if eligible and verification is not None
        else {}
    )
    return {
        "case_id": candidate_id,
        "panel_count": 1,
        "split": curated.get("split"),
        "data_path": table["path"],
        "sheet": table.get("sheet_name"),
        "data_sha256": table["sha256"],
        "panel_id": mapping.panel_id,
        "user_goal": curated.get("user_goal"),
        "chart_family": curated.get("chart_family"),
        "intent": curated.get("intent"),
        "evaluation_expectation": curated.get("evaluation_expectation"),
        "curation_status": curation_status,
        "eligible_for_experiment": eligible,
        "eligibility_reasons": [] if eligible else ["curation-not-verified"],
    }


def _iter_source_tables(
    article_dir: Path,
    *,
    article_id: str,
    extraction_root: Path,
    max_zip_files: int,
    max_zip_uncompressed_bytes: int,
) -> tuple[list[tuple[Path, Path | None, str | None]], list[dict[str, Any]], dict[str, int]]:
    source_dir = article_dir / "source_data"
    tables: list[tuple[Path, Path | None, str | None]] = []
    ambiguous: list[dict[str, Any]] = []
    counts = {
        "archives_processed": 0,
        "archives_rejected": 0,
        "direct_tables": 0,
        "extracted_tables": 0,
        "resource_files_skipped": 0,
    }
    if not source_dir.is_dir():
        return tables, ambiguous, counts
    for path in sorted(source_dir.iterdir(), key=lambda item: item.name.casefold()):
        if path.is_symlink():
            ambiguous.append(
                {
                    "article_id": article_id,
                    "source_path": str(path.resolve(strict=False)),
                    "reasons": ["source-data-symlink-rejected"],
                }
            )
            continue
        if not path.is_file():
            continue
        if _is_resource_fork_name(path.name):
            counts["resource_files_skipped"] += 1
            continue
        suffix = path.suffix.casefold()
        if suffix in TABLE_SUFFIXES:
            counts["direct_tables"] += 1
            tables.append((path, None, None))
            continue
        if suffix != ".zip":
            continue
        archive_hash = sha256_file(path)
        destination = (
            extraction_root
            / article_id
            / f"{path.stem}-{archive_hash[:12]}"
        )
        try:
            extracted = safe_extract_tables(
                path,
                destination,
                max_files=max_zip_files,
                max_uncompressed_bytes=max_zip_uncompressed_bytes,
            )
        except ArchiveSafetyError as exc:
            counts["archives_rejected"] += 1
            ambiguous.append(
                {
                    "article_id": article_id,
                    "source_path": str(path.resolve()),
                    "source_sha256": archive_hash,
                    "source_size_bytes": path.stat().st_size,
                    "reasons": [exc.code],
                    "detail": exc.detail,
                }
            )
            continue
        counts["archives_processed"] += 1
        if not extracted:
            ambiguous.append(
                {
                    "article_id": article_id,
                    "source_path": str(path.resolve()),
                    "source_sha256": archive_hash,
                    "source_size_bytes": path.stat().st_size,
                    "reasons": ["archive-no-csv-or-xlsx"],
                }
            )
            continue
        for table in extracted:
            counts["extracted_tables"] += 1
            tables.append((table.path, path, table.member_name))
    return tables, ambiguous, counts


def build_cases(
    *,
    corpus_manifest: str | Path,
    content_root: str | Path,
    output_root: str | Path,
    evidence_file: str | Path | None = None,
    max_zip_files: int = DEFAULT_MAX_ZIP_FILES,
    max_zip_uncompressed_bytes: int = DEFAULT_MAX_ZIP_UNCOMPRESSED_BYTES,
    max_xlsx_sheets: int = DEFAULT_MAX_XLSX_SHEETS,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    if max_xlsx_sheets < 1:
        raise ValueError("max_xlsx_sheets must be positive")
    manifest_path = Path(corpus_manifest)
    content = Path(content_root).resolve()
    output = Path(output_root).resolve()
    output.mkdir(parents=True, exist_ok=True)
    extraction_root = output / "extracted"
    shutil.rmtree(extraction_root, ignore_errors=True)
    extraction_root.mkdir(parents=True)
    manifest_hash = sha256_file(manifest_path)
    evidence, invalid_evidence = _load_evidence(evidence_file)
    evidence_file_hash = sha256_file(evidence_file) if evidence_file else None
    records = sorted(
        _read_manifest(manifest_path),
        key=lambda record: normalize_doi(record.get("doi")) or "",
    )

    candidates: list[dict[str, Any]] = []
    ambiguous: list[dict[str, Any]] = []
    totals = {
        "archives_processed": 0,
        "archives_rejected": 0,
        "direct_tables": 0,
        "extracted_tables": 0,
        "resource_files_skipped": 0,
        "xlsx_workbooks": 0,
        "xlsx_sheets_inspected": 0,
        "xlsx_resource_sheets_skipped": 0,
        "xlsx_corrupt": 0,
        "xlsx_sheet_limit_exceeded": 0,
    }
    eligible_articles = 0
    skipped_articles = 0

    for record in records:
        doi = normalize_doi(record.get("doi"))
        article_id = _article_id(record)
        if not doi or not article_id:
            skipped_articles += 1
            ambiguous.append(
                {
                    "doi": doi,
                    "article_id": article_id,
                    "reasons": ["article-identity-missing"],
                    "curation_status": "unverified",
                    "eligible_for_experiment": False,
                }
            )
            continue
        license_record = record.get("license") or {}
        license_url = license_record.get("normalized_url") or license_record.get("url")
        license_source = license_record.get("source") or record.get("license_source")
        license_evidence = license_record.get("evidence") or record.get(
            "license_evidence"
        )
        if (
            not record.get("download_eligible")
            or not normalize_cc_by_url(license_url)
            or license_source not in {"crossref", "article_metadata"}
            or not license_evidence
        ):
            skipped_articles += 1
            ambiguous.append(
                {
                    "doi": doi,
                    "article_id": article_id,
                    "reasons": ["article-not-cc-by-download-eligible"],
                    "curation_status": "unverified",
                    "eligible_for_experiment": False,
                }
            )
            continue
        eligible_articles += 1
        article_dir = content / article_id
        provenance_path = article_dir / "meta" / "provenance.json"
        provenance_hash, article_reasons = _article_provenance(
            provenance_path,
            doi=doi,
            license_url=str(license_url),
        )
        tables, archive_ambiguities, counts = _iter_source_tables(
            article_dir,
            article_id=article_id,
            extraction_root=extraction_root,
            max_zip_files=max_zip_files,
            max_zip_uncompressed_bytes=max_zip_uncompressed_bytes,
        )
        for key, value in counts.items():
            totals[key] += value
        for item in archive_ambiguities:
            ambiguous.append(
                {
                    "schema_version": SCHEMA_VERSION,
                    "doi": doi,
                    "corpus_manifest_sha256": manifest_hash,
                    "provenance_manifest_sha256": provenance_hash,
                    "curation_status": "unverified",
                    "eligible_for_experiment": False,
                    **item,
                }
            )

        for table_path, archive_path, archive_member in sorted(
            tables,
            key=lambda item: (
                str(item[1] or ""),
                str(item[2] or item[0].name).casefold(),
            ),
        ):
            base_table = _source_table_descriptor(
                table_path,
                content_root=content,
                output_root=output,
                archive_path=archive_path,
                archive_member=archive_member,
                sheet_name=None,
            )
            attempts, workbook_metrics = _mapping_attempts(
                table_path,
                mapping_name=archive_member or table_path.name,
                max_xlsx_sheets=max_xlsx_sheets,
            )
            for key, value in workbook_metrics.items():
                totals[key] += value
            for attempt in attempts:
                table = {**base_table, "sheet_name": attempt.sheet_name}
                mapping = attempt.mapping
                reasons = list(article_reasons) + list(attempt.reasons)
                figure = None
                caption = None
                if mapping:
                    figure, caption, asset_reasons = _figure_assets(
                        article_dir,
                        mapping.figure_no,
                    )
                    reasons.extend(asset_reasons)
                if reasons or not mapping:
                    ambiguous.append(
                        {
                            "schema_version": SCHEMA_VERSION,
                            "doi": doi,
                            "article_id": article_id,
                            "figure_no": mapping.figure_no if mapping else None,
                            "panel_ids": [mapping.panel_id] if mapping else [],
                            "panel_qualifier": mapping.qualifier if mapping else None,
                            "source_table": table,
                            "figure": figure,
                            "caption": caption,
                            "license": license_record,
                            "license_evidence_sha256": _canonical_json_sha256(
                                license_record
                            ),
                            "corpus_manifest_sha256": manifest_hash,
                            "provenance_manifest_sha256": provenance_hash,
                            "reasons": sorted(set(reasons)),
                            "detail": attempt.detail,
                            "curation_status": "unverified",
                            "eligible_for_experiment": False,
                        }
                    )
                    continue

                candidate_id = _candidate_id(
                    article_id=article_id,
                    doi=doi,
                    mapping=mapping,
                    table=table,
                )
                verification = evidence.get(candidate_id)
                curation_status = "verified" if verification else "unverified"
                skeleton = _case_skeleton(
                    candidate_id=candidate_id,
                    mapping=mapping,
                    table=table,
                    curation_status=curation_status,
                    verification=verification,
                )
                candidates.append(
                    {
                        "schema_version": SCHEMA_VERSION,
                        "candidate_id": candidate_id,
                        "doi": doi,
                        "figure_no": mapping.figure_no,
                        "panel_ids": [mapping.panel_id],
                        "panel_qualifier": mapping.qualifier,
                        "source_table": table,
                        "figure": figure,
                        "caption": caption,
                        "license": license_record,
                        "license_evidence_sha256": _canonical_json_sha256(
                            license_record
                        ),
                        "corpus_manifest_sha256": manifest_hash,
                        "provenance_manifest_sha256": provenance_hash,
                        "curation_status": curation_status,
                        "verification_evidence": verification,
                        "verification_evidence_file_sha256": (
                            evidence_file_hash if verification else None
                        ),
                        "eligible_for_experiment": (
                            curation_status == "verified"
                        ),
                        "eligibility_reasons": (
                            [] if curation_status == "verified"
                            else ["curation-not-verified"]
                        ),
                        "experiment_case": skeleton,
                    }
                )

    candidates.sort(key=lambda item: item["candidate_id"])
    ambiguous.sort(
        key=lambda item: (
            str(item.get("doi")),
            str((item.get("source_table") or {}).get("path") or item.get("source_path")),
            str((item.get("source_table") or {}).get("sheet_name") or ""),
            ",".join(item.get("reasons") or []),
        )
    )
    verified = sum(
        candidate["curation_status"] == "verified" for candidate in candidates
    )
    summary = {
        "schema_version": SCHEMA_VERSION,
        "corpus_manifest": str(manifest_path.resolve()),
        "corpus_manifest_sha256": manifest_hash,
        "content_root": str(content),
        "articles_total": len(records),
        "articles_eligible": eligible_articles,
        "articles_skipped": skipped_articles,
        **totals,
        "candidates": len(candidates),
        "ambiguous": len(ambiguous),
        "verified": verified,
        "unverified": len(candidates) - verified,
        "eligible_for_experiment": verified,
        "invalid_evidence_records": invalid_evidence,
        "evidence_file_sha256": evidence_file_hash,
        "max_zip_files": max_zip_files,
        "max_zip_uncompressed_bytes": max_zip_uncompressed_bytes,
        "max_xlsx_sheets": max_xlsx_sheets,
        "llm_calls": 0,
    }
    return candidates, ambiguous, summary


def _write_jsonl(path: Path, records: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(
                json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n"
            )


def write_case_outputs(
    output_root: str | Path,
    candidates: list[dict[str, Any]],
    ambiguous: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    output = Path(output_root)
    output.mkdir(parents=True, exist_ok=True)
    _write_jsonl(output / "candidates.jsonl", candidates)
    _write_jsonl(output / "ambiguous.jsonl", ambiguous)
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
