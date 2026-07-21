from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import stat
import zipfile

import pytest
from openpyxl import Workbook

from nature_download.corpus import cases as cases_module
from nature_download.corpus.cases import (
    ArchiveSafetyError,
    build_cases,
    parse_figure_panel,
    safe_extract_tables,
    write_case_outputs,
)
from nature_download.corpus.provenance import sha256_file


ARTICLE_ID = "s41467-024-24680-2"
DOI = f"10.1038/{ARTICLE_ID}"


def write_zip(path: Path, members: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as handle:
        for name, payload in members.items():
            handle.writestr(name, payload)


def workbook_bytes(sheet_names: list[str]) -> bytes:
    workbook = Workbook()
    workbook.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        workbook.create_sheet(title=name)
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


def test_safe_zip_rejects_zip_slip_and_absolute_paths(workdir: Path) -> None:
    for name in ("../escape.csv", "/absolute.csv", "C:\\absolute.csv"):
        archive = workdir / f"bad-{len(name)}.zip"
        write_zip(archive, {name: b"x,y\n1,2\n"})
        destination = workdir / f"extract-{len(name)}"
        with pytest.raises(
            ArchiveSafetyError,
            match="zip-slip|zip-absolute-path",
        ):
            safe_extract_tables(archive, destination)
        assert not (workdir / "escape.csv").exists()
        assert not destination.exists()


def test_safe_zip_rejects_symlink_member(workdir: Path) -> None:
    archive = workdir / "symlink.zip"
    info = zipfile.ZipInfo("Figure2B.csv")
    info.create_system = 3
    info.external_attr = (stat.S_IFLNK | 0o777) << 16
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr(info, "../../outside")
    with pytest.raises(ArchiveSafetyError, match="zip-symlink"):
        safe_extract_tables(archive, workdir / "extract")


def test_safe_zip_enforces_expanded_size_and_file_count(workdir: Path) -> None:
    archive = workdir / "bomb.zip"
    write_zip(
        archive,
        {
            "Figure2B.csv": b"A" * 200,
            "Figure2C.csv": b"B" * 200,
        },
    )
    with pytest.raises(
        ArchiveSafetyError,
        match="zip-uncompressed-size-limit",
    ):
        safe_extract_tables(
            archive,
            workdir / "extract-size",
            max_uncompressed_bytes=100,
        )
    with pytest.raises(ArchiveSafetyError, match="zip-file-count-limit"):
        safe_extract_tables(
            archive,
            workdir / "extract-count",
            max_files=1,
        )


def test_safe_zip_extracts_only_csv_and_xlsx(workdir: Path) -> None:
    archive = workdir / "tables.zip"
    write_zip(
        archive,
        {
            "nested/Figure2B.csv": b"x,y\n1,2\n",
            "Figure3J-inset.xlsx": b"xlsx bytes",
            "__MACOSX/._Figure4A.xlsx": b"resource fork",
            "script.py": b"raise RuntimeError('must not execute')",
        },
    )
    extracted = safe_extract_tables(archive, workdir / "extract")
    assert [item.member_name for item in extracted] == [
        "Figure3J-inset.xlsx",
        "nested/Figure2B.csv",
    ]
    assert not (workdir / "extract" / "script.py").exists()


def test_xlsx_sheet_inspection_uses_read_only_mode(
    monkeypatch: pytest.MonkeyPatch,
    workdir: Path,
) -> None:
    calls: list[dict] = []

    class FakeWorkbook:
        sheetnames = ["Fig. 1a"]

        def close(self) -> None:
            return None

    def fake_load_workbook(**kwargs):
        calls.append(kwargs)
        return FakeWorkbook()

    monkeypatch.setattr(cases_module, "load_workbook", fake_load_workbook)
    workbook = workdir / "generic.xlsx"
    workbook.write_bytes(b"unused by fake loader")
    sheets = cases_module._read_xlsx_sheet_names(workbook, max_sheets=256)
    assert sheets.names == ("Fig. 1a",)
    assert calls == [
        {
            "filename": workbook,
            "read_only": True,
            "data_only": False,
            "keep_links": False,
        }
    ]


@pytest.mark.parametrize(
    ("filename", "figure_no", "panel_id", "qualifier"),
    [
        ("Figure2B.csv", 2, "b", None),
        ("Figure_3C.xlsx", 3, "c", None),
        ("Figure3J-inset.csv", 3, "j", "inset"),
        ("Fig4K-inset.csv", 4, "k", "inset"),
        ("Fig. 1a", 1, "a", None),
        ("Figure 3C", 3, "c", None),
    ],
)
def test_figure_panel_mapping(
    filename: str,
    figure_no: int,
    panel_id: str,
    qualifier: str | None,
) -> None:
    mapping, reasons = parse_figure_panel(filename)
    assert reasons == []
    assert mapping is not None
    assert (mapping.figure_no, mapping.panel_id, mapping.qualifier) == (
        figure_no,
        panel_id,
        qualifier,
    )


@pytest.mark.parametrize(
    "name",
    [
        "Figure1",
        "Fig 2",
        "Fig2G-H",
        "Figure 5b,c",
        "Figure2BC",
        "Supplementary Figure 2B",
        "Sup Fig 1A",
        "Sup_Fig_1A",
        "Supp_Figure_1A",
        "Supplementary_Figure_2B",
    ],
)
def test_non_unique_or_missing_panel_reference_is_ambiguous(name: str) -> None:
    mapping, reasons = parse_figure_panel(name)
    assert mapping is None
    assert reasons


def make_case_fixture(workdir: Path) -> tuple[Path, Path]:
    content = workdir / "content"
    article = content / ARTICLE_ID
    figures = article / "figures"
    source_data = article / "source_data"
    meta = article / "meta"
    figures.mkdir(parents=True)
    source_data.mkdir()
    meta.mkdir()
    for figure_no in (1, 2, 3):
        (figures / f"fig_{figure_no:03d}.png").write_bytes(
            f"figure {figure_no}".encode()
        )
        (figures / f"fig_{figure_no:03d}.txt").write_text(
            f"caption {figure_no}",
            encoding="utf-8",
        )
    (source_data / "Figure2B.csv").write_text("x,y\n1,2\n", encoding="utf-8")
    (source_data / "Figure2C.xlsx").write_bytes(
        workbook_bytes(["Measurements"])
    )
    (source_data / "unmapped.csv").write_text("x,y\n3,4\n", encoding="utf-8")
    (source_data / "corrupt.xlsx").write_bytes(b"not a workbook")
    (source_data / "._ignored.xlsx").write_bytes(b"resource fork")
    write_zip(
        source_data / "source.zip",
        {
            "tables/source.xlsx": workbook_bytes(
                [
                    "Fig. 1a",
                    "Figure 3C",
                    "Figure3J-inset",
                    "Figure1",
                    "Fig2G-H",
                    "Figure 5b,c",
                    "._resource",
                ]
            )
        },
    )
    file_entries = []
    for path in sorted(
        [
            item
            for folder in (figures, source_data)
            for item in folder.iterdir()
            if item.is_file()
        ]
    ):
        kind = (
            "source_data"
            if path.parent == source_data
            else "caption"
            if path.suffix == ".txt"
            else "figure"
        )
        entry = {
            "kind": kind,
            "path": path.relative_to(article).as_posix(),
            "sha256": sha256_file(path),
            "size_bytes": path.stat().st_size,
            "source_url": f"https://example.test/{path.name}",
            "download_status": "downloaded",
            "rejection_reason": None,
        }
        if kind == "source_data":
            entry.update(
                {
                    "source_data_origin": "supplementary_information",
                    "verification_status": "source-provided",
                    "verification_evidence": "test fixture",
                }
            )
        file_entries.append(entry)
    (meta / "provenance.json").write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "doi": DOI,
                "journal": "Nature Communications",
                "article_url": (
                    f"https://www.nature.com/articles/{ARTICLE_ID}"
                ),
                "journal_allowed": True,
                "require_cc_by": True,
                "download_eligible": True,
                "download_status": "downloaded",
                "rejection_reasons": [],
                "source_data_origin": "supplementary_information",
                "license_source": "crossref",
                "license_evidence": {
                    "URL": "https://creativecommons.org/licenses/by/4.0/",
                    "content-version": "vor",
                },
                "license": {
                    "normalized_url": (
                        "https://creativecommons.org/licenses/by/4.0/"
                    ),
                    "source": "crossref",
                    "content_version": "vor",
                },
                "files": file_entries,
            }
        ),
        encoding="utf-8",
    )
    manifest = workdir / "corpus_manifest.jsonl"
    manifest.write_text(
        json.dumps(
            {
                "doi": DOI,
                "article_url": f"https://www.nature.com/articles/{ARTICLE_ID}",
                "download_eligible": True,
                "license": {
                    "normalized_url": (
                        "https://creativecommons.org/licenses/by/4.0/"
                    ),
                    "version": "4.0",
                    "content_version": "vor",
                    "source": "crossref",
                    "evidence": {
                        "URL": "https://creativecommons.org/licenses/by/4.0/",
                        "content-version": "vor",
                    },
                },
            }
        )
        + "\n",
        encoding="utf-8",
    )
    return manifest, content


def remove_provenance_file(content: Path, relative_path: str) -> None:
    path = content / ARTICLE_ID / "meta" / "provenance.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    manifest["files"] = [
        entry
        for entry in manifest["files"]
        if entry.get("path") != relative_path
    ]
    path.write_text(json.dumps(manifest), encoding="utf-8")


def test_case_builder_checksums_ambiguity_and_unverified_gate(workdir: Path) -> None:
    manifest, content = make_case_fixture(workdir)
    output = workdir / "cases"
    candidates, ambiguous, summary = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=output,
    )
    assert summary["candidates"] == 5
    assert summary["ambiguous"] == 5
    assert summary["direct_tables"] == 4
    assert summary["extracted_tables"] == 1
    assert summary["resource_files_skipped"] == 1
    assert summary["xlsx_workbooks"] == 3
    assert summary["xlsx_sheets_inspected"] == 8
    assert summary["xlsx_resource_sheets_skipped"] == 1
    assert summary["xlsx_corrupt"] == 1
    assert summary["verified"] == 0
    assert summary["eligible_for_experiment"] == 0
    assert summary["llm_calls"] == 0


    assert {item["figure_no"] for item in candidates} == {1, 2, 3}
    assert {item["panel_ids"][0] for item in candidates} == {"a", "b", "c", "j"}
    assert all(item["curation_status"] == "unverified" for item in candidates)
    assert all(item["eligible_for_experiment"] is False for item in candidates)
    assert all(
        item["eligibility_reasons"] == ["curation-not-verified"]
        for item in candidates
    )
    assert all(
        item["experiment_case"]["eligible_for_experiment"] is False
        for item in candidates
    )
    assert all(item["experiment_case"]["split"] is None for item in candidates)
    assert all(
        item["experiment_case"]["user_goal"] is None
        for item in candidates
    )
    direct = next(
        item
        for item in candidates
        if item["figure_no"] == 2
        and item["panel_ids"] == ["b"]
        and item["source_table"]["sheet_name"] is None
    )
    assert direct["source_table"]["sha256"] == sha256_file(
        content / ARTICLE_ID / "source_data" / "Figure2B.csv"
    )
    assert direct["figure"]["sha256"] == sha256_file(
        content / ARTICLE_ID / "figures" / "fig_002.png"
    )
    assert direct["caption"]["sha256"] == sha256_file(
        content / ARTICLE_ID / "figures" / "fig_002.txt"
    )
    sheet_candidates = [
        item for item in candidates if item["source_table"]["sheet_name"]
    ]
    assert {item["source_table"]["sheet_name"] for item in sheet_candidates} == {
        "Fig. 1a",
        "Figure 3C",
        "Figure3J-inset",
    }
    assert all(
        item["experiment_case"]["data_path"] == item["source_table"]["path"]
        and item["experiment_case"]["sheet"]
        == item["source_table"]["sheet_name"]
        for item in sheet_candidates
    )
    assert len({item["candidate_id"] for item in sheet_candidates}) == 3
    assert len({item["source_table"]["sha256"] for item in sheet_candidates}) == 1
    assert all(item["eligible_for_experiment"] is False for item in ambiguous)
    assert any(
        item["reasons"] == ["xlsx-workbook-invalid"] for item in ambiguous
    )
    assert {
        item["source_table"]["sheet_name"]
        for item in ambiguous
        if (item.get("source_table") or {}).get("sheet_name")
    } == {"Figure1", "Fig2G-H", "Figure 5b,c"}


def test_case_builder_rejects_untracked_direct_source_data(workdir: Path) -> None:
    manifest, content = make_case_fixture(workdir)
    untracked = content / ARTICLE_ID / "source_data" / "Figure1B.csv"
    untracked.write_text("x,y\n1,2\n", encoding="utf-8")

    candidates, ambiguous, _ = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=workdir / "cases",
    )

    assert all(item["source_table"]["path"] != str(untracked) for item in candidates)
    assert any(
        item["source_table"]["path"] == str(untracked)
        and "source-table-not-in-article-provenance" in item["reasons"]
        for item in ambiguous
        if item.get("source_table")
    )


def test_case_builder_requires_selected_archive_in_provenance(workdir: Path) -> None:
    manifest, content = make_case_fixture(workdir)
    remove_provenance_file(content, "source_data/source.zip")

    candidates, ambiguous, _ = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=workdir / "cases",
    )

    assert all(item["source_table"]["archive"] is None for item in candidates)
    assert any(
        "source-table-not-in-article-provenance" in item.get("reasons", [])
        and (item.get("source_table") or {}).get("archive") is not None
        for item in ambiguous
    )


@pytest.mark.parametrize(
    ("relative_path", "reason"),
    [
        ("figures/fig_002.png", "figure-not-in-article-provenance"),
        ("figures/fig_002.txt", "caption-not-in-article-provenance"),
    ],
)
def test_case_builder_requires_selected_figure_assets_in_provenance(
    workdir: Path,
    relative_path: str,
    reason: str,
) -> None:
    manifest, content = make_case_fixture(workdir)
    remove_provenance_file(content, relative_path)

    candidates, ambiguous, _ = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=workdir / "cases",
    )

    assert all(item["figure_no"] != 2 for item in candidates)
    assert any(
        item.get("figure_no") == 2 and reason in item.get("reasons", [])
        for item in ambiguous
    )


def test_case_builder_malformed_provenance_files_fails_closed(
    workdir: Path,
) -> None:
    manifest, content = make_case_fixture(workdir)
    path = content / ARTICLE_ID / "meta" / "provenance.json"
    provenance = json.loads(path.read_text(encoding="utf-8"))
    provenance["files"] = {"not": "a-list"}
    path.write_text(json.dumps(provenance), encoding="utf-8")

    candidates, ambiguous, _ = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=workdir / "cases",
    )

    assert candidates == []
    assert any(
        "article-provenance:files-invalid" in item.get("reasons", [])
        for item in ambiguous
    )


def test_case_builder_rejects_file_changed_after_article_manifest(
    workdir: Path,
) -> None:
    manifest, content = make_case_fixture(workdir)
    changed = content / ARTICLE_ID / "source_data" / "Figure2B.csv"
    changed.write_text("x,y\n1,999\n", encoding="utf-8")

    candidates, ambiguous, summary = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=workdir / "cases",
    )

    assert candidates == []
    assert summary["verified"] == 0
    assert any(
        any(
            reason.startswith(
                "article-provenance:checksum-mismatch:source_data/Figure2B.csv"
            )
            for reason in item.get("reasons") or []
        )
        for item in ambiguous
    )


def test_xlsx_sheet_limit_is_recorded_without_loading_cells(workdir: Path) -> None:
    manifest, content = make_case_fixture(workdir)
    candidates, ambiguous, summary = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=workdir / "cases",
        max_xlsx_sheets=2,
    )
    assert summary["max_xlsx_sheets"] == 2
    assert summary["xlsx_sheet_limit_exceeded"] == 1
    assert any(
        item["reasons"] == ["xlsx-sheet-count-limit"]
        for item in ambiguous
    )
    assert all(
        item["source_table"]["sheet_name"] is None
        for item in ambiguous
        if item["reasons"] == ["xlsx-sheet-count-limit"]
    )
    assert all(item["eligible_for_experiment"] is False for item in candidates)


def test_only_explicit_valid_evidence_can_verify(workdir: Path) -> None:
    manifest, content = make_case_fixture(workdir)
    output = workdir / "cases"
    initial, _, _ = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=output,
    )
    candidate_id = next(
        item["candidate_id"]
        for item in initial
        if item["panel_ids"] == ["b"]
        and item["source_table"]["sheet_name"] is None
    )
    invalid_evidence = workdir / "invalid-evidence.json"
    invalid_evidence.write_text(
        json.dumps(
            {
                candidate_id: {
                    "status": "verified",
                    "evidence_type": "automatic",
                    "evidence_ref": "auto",
                    "reviewer_or_source": "script",
                }
            }
        ),
        encoding="utf-8",
    )
    candidates, _, summary = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=output,
        evidence_file=invalid_evidence,
    )
    assert summary["invalid_evidence_records"] == 1
    assert all(not item["eligible_for_experiment"] for item in candidates)

    evidence = workdir / "evidence.json"
    evidence.write_text(
        json.dumps(
            {
                candidate_id: {
                    "status": "verified",
                    "evidence_type": "human_review",
                    "evidence_ref": "review-form-001",
                    "reviewer_or_source": "curator-01",
                    "experiment_case": {
                        "user_goal": "Plot the verified source series.",
                        "chart_family": "line",
                        "intent": {"x": "x", "y": "y"},
                        "evaluation_expectation": {
                            "schema_version": "1.1.0",
                            "panels": [
                                {
                                    "panel_id": "b",
                                    "series": [
                                        {
                                            "series_id": "verified",
                                            "kind": "line",
                                            "x": "x",
                                            "y": "y"
                                        }
                                    ]
                                }
                            ],
                            "panel_groups": []
                        }
                    }
                }
            }
        ),
        encoding="utf-8",
    )
    candidates, _, summary = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=output,
        evidence_file=evidence,
    )
    verified = [
        item for item in candidates if item["candidate_id"] == candidate_id
    ][0]
    assert verified["curation_status"] == "verified"
    assert verified["eligible_for_experiment"] is True
    assert verified["eligibility_reasons"] == []
    assert verified["experiment_case"]["eligible_for_experiment"] is True
    assert summary["verified"] == 1


def test_case_outputs_are_deterministic(workdir: Path) -> None:
    manifest, content = make_case_fixture(workdir)
    output = workdir / "cases"
    first = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=output,
    )
    write_case_outputs(output, *first)
    first_bytes = {
        name: (output / name).read_bytes()
        for name in ("candidates.jsonl", "ambiguous.jsonl", "summary.json")
    }
    second = build_cases(
        corpus_manifest=manifest,
        content_root=content,
        output_root=output,
    )
    write_case_outputs(output, *second)
    assert first == second
    assert first_bytes == {
        name: (output / name).read_bytes()
        for name in ("candidates.jsonl", "ambiguous.jsonl", "summary.json")
    }
